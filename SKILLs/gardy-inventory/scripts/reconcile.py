#!/usr/bin/env python3
"""Reconcile two inventory snapshots without changing either source file."""

import argparse
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path


def columns(value):
    result = [part.strip() for part in value.split(',')]
    if not result or any(not part for part in result) or len(set(result)) != len(result):
        raise ValueError('匹配字段必须是互不重复的非空列名，用英文逗号分隔。')
    return result


def cell_text(value, number_format=None):
    if value is None:
        return ''
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, int) and number_format and re.fullmatch(r'0+', number_format):
        return str(value).zfill(len(number_format))
    return str(value).strip()


def read_table(path, sheet_name, warnings, identifier_columns):
    suffix = path.suffix.lower()
    if suffix == '.csv':
        if sheet_name:
            raise ValueError(f'{path.name} 是 CSV 文件，不能指定工作表。')
        with path.open('r', encoding='utf-8-sig', newline='') as handle:
            reader = csv.reader(handle)
            try:
                header = next(reader)
            except StopIteration:
                raise ValueError(f'{path.name} 是空文件。') from None
            rows = [(line, values) for line, values in enumerate(reader, start=2)]
        return [cell_text(value) for value in header], rows

    if suffix not in ('.xlsx', '.xlsm'):
        raise ValueError(f'{path.name} 的格式不受支持；请使用 XLSX、XLSM 或 UTF-8 CSV。')
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError('读取 Excel 需要 openpyxl。请在当前 Python 环境安装后重试：python3 -m pip install openpyxl') from None

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == 'visible']
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f'{path.name} 没有工作表“{sheet_name}”。')
            sheet = workbook[sheet_name]
        elif len(visible) == 1:
            sheet = visible[0]
        else:
            raise ValueError(f'{path.name} 有多个工作表，请明确指定：{", ".join(sheet.title for sheet in visible)}')
        iterator = sheet.iter_rows()
        try:
            header_cells = next(iterator)
        except StopIteration:
            raise ValueError(f'{path.name} 的工作表是空的。') from None
        header = [cell_text(cell.value) for cell in header_cells]
        identifier_indexes = [header.index(name) for name in identifier_columns if name in header]
        rows = []
        numeric_identifier_rows = set()
        for row in iterator:
            values = [cell_text(cell.value, cell.number_format) for cell in row]
            rows.append((row[0].row, values))
            for index in identifier_indexes:
                if index >= len(row):
                    continue
                cell = row[index]
                if isinstance(cell.value, (int, float)) and cell.number_format == 'General':
                    numeric_identifier_rows.add(cell.row)
        if numeric_identifier_rows:
            warnings.append(f'{path.name} 有数字单元格采用通用格式；若它们是带前导零的编号，原始 Excel 可能已丢失前导零。')
        return header, rows
    finally:
        workbook.close()


def parse_quantity(value, source, line):
    cleaned = value.replace(',', '').strip()
    if not cleaned:
        raise ValueError(f'{source} 第 {line} 行数量为空。')
    try:
        quantity = Decimal(cleaned)
    except InvalidOperation:
        raise ValueError(f'{source} 第 {line} 行数量“{value}”不是数字。') from None
    if not quantity.is_finite():
        raise ValueError(f'{source} 第 {line} 行数量不是有限数字。')
    return quantity


def load_source(path, sheet, key_columns, quantity_column, warnings):
    header, raw_rows = read_table(path, sheet, warnings, key_columns)
    if not header or any(not field for field in header) or len(set(header)) != len(header):
        raise ValueError(f'{path.name} 的表头有空列名或重复列名。')
    required = [*key_columns, quantity_column]
    missing = [field for field in required if field not in header]
    if missing:
        raise ValueError(f'{path.name} 缺少字段：{", ".join(missing)}；现有字段：{", ".join(header)}')

    key_indexes = [header.index(field) for field in key_columns]
    quantity_index = header.index(quantity_column)
    totals = defaultdict(Decimal)
    counts = Counter()
    duplicates = defaultdict(list)
    total = Decimal(0)
    row_count = 0
    for line, raw in raw_rows:
        values = [cell_text(value) for value in (raw + [''] * len(header))[:len(header)]]
        if not any(value.strip() for value in values):
            continue
        if len(raw) > len(header) and any(value.strip() for value in raw[len(header):]):
            raise ValueError(f'{path.name} 第 {line} 行的单元格数量超过表头。')
        key = tuple(values[index] for index in key_indexes)
        if any(not value for value in key):
            raise ValueError(f'{path.name} 第 {line} 行的匹配字段为空。')
        quantity = parse_quantity(values[quantity_index], path.name, line)
        totals[key] += quantity
        counts[key] += 1
        duplicates[tuple(values)].append(line)
        total += quantity
        row_count += 1
    return {
        'header': header,
        'totals': totals,
        'counts': counts,
        'duplicates': duplicates,
        'row_count': row_count,
        'total': total,
    }


def number(value):
    return format(value, 'f')


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def reconcile(args):
    ledger_path = Path(args.ledger).expanduser().resolve()
    count_path = Path(args.count).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    if output_path.suffix.lower() != '.csv':
        raise ValueError('差异清单输出路径必须以 .csv 结尾。')
    duplicates_path = output_path.with_name(output_path.stem + '_完全重复行.csv')
    for source in (ledger_path, count_path):
        if not source.is_file():
            raise ValueError(f'找不到源文件：{source}')
    if ledger_path == count_path:
        raise ValueError('账面与盘点必须使用不同的源文件。')
    if output_path in (ledger_path, count_path) or duplicates_path in (ledger_path, count_path):
        raise ValueError('输出路径不能覆盖源文件。')

    ledger_keys = columns(args.ledger_keys or args.keys)
    count_keys = columns(args.count_keys or args.keys)
    if len(ledger_keys) != len(count_keys):
        raise ValueError('账面与盘点的匹配字段数量必须相同。')
    warnings = []
    ledger = load_source(ledger_path, args.ledger_sheet, ledger_keys, args.ledger_quantity, warnings)
    count = load_source(count_path, args.count_sheet, count_keys, args.count_quantity, warnings)

    matched = 0
    ledger_only = 0
    count_only = 0
    rows = []
    all_keys = sorted(set(ledger['totals']) | set(count['totals']))
    for key in all_keys:
        in_ledger = key in ledger['totals']
        in_count = key in count['totals']
        if in_ledger and in_count:
            status = '已匹配'
            note = ''
            matched += 1
        elif in_ledger:
            status = '盘点无对应记录'
            note = '差异计算以零占位；不代表确认零库存，需人工核对'
            ledger_only += 1
        else:
            status = '账面无对应记录'
            note = '差异计算以零占位；不代表确认零库存，需人工核对'
            count_only += 1
        ledger_quantity = ledger['totals'].get(key, Decimal(0))
        count_quantity = count['totals'].get(key, Decimal(0))
        rows.append([
            *key,
            number(ledger_quantity) if in_ledger else '',
            number(count_quantity) if in_count else '',
            number(count_quantity - ledger_quantity),
            status,
            note,
            ledger['counts'][key],
            count['counts'][key],
        ])

    write_csv(output_path, [
        *ledger_keys, '账面数量', '盘点数量', '差异_盘点减账面', '状态',
        '说明', '账面原始行数', '盘点原始行数',
    ], rows)
    duplicate_rows = []
    for source_name, source in ((ledger_path.name, ledger), (count_path.name, count)):
        for signature, lines in source['duplicates'].items():
            if len(lines) > 1:
                duplicate_rows.append([source_name, len(lines), ','.join(map(str, lines)), json.dumps(dict(zip(source['header'], signature)), ensure_ascii=False)])
    write_csv(duplicates_path, ['来源文件', '出现次数', '原始行号', '原始行内容_JSON'], duplicate_rows)

    return {
        '账面原始行数': ledger['row_count'],
        '盘点原始行数': count['row_count'],
        '账面合计': number(ledger['total']),
        '盘点合计': number(count['total']),
        '总差异_盘点减账面': number(count['total'] - ledger['total']),
        '匹配组数': matched,
        '盘点无对应记录组数': ledger_only,
        '账面无对应记录组数': count_only,
        '完全重复组数': len(duplicate_rows),
        '差异清单': str(output_path),
        '重复行清单': str(duplicates_path),
        '重要说明': '无对应记录不等于确认零库存；差异仅供人工核对，不能直接调账或认定责任。',
        '警告': warnings,
    }


def main():
    parser = argparse.ArgumentParser(description='嘉迪库存账面与盘点文件核对（输出中文 CSV 与 JSON 摘要）')
    parser.add_argument('--ledger', required=True, help='账面 XLSX、XLSM 或 UTF-8 CSV 路径')
    parser.add_argument('--count', required=True, help='盘点 XLSX、XLSM 或 UTF-8 CSV 路径')
    parser.add_argument('--keys', required=True, help='两边共有的匹配字段，英文逗号分隔')
    parser.add_argument('--ledger-keys', help='账面匹配字段，与 --keys 顺序对应')
    parser.add_argument('--count-keys', help='盘点匹配字段，与 --keys 顺序对应')
    parser.add_argument('--ledger-quantity', required=True, help='账面数量列名')
    parser.add_argument('--count-quantity', required=True, help='盘点数量列名')
    parser.add_argument('--ledger-sheet', help='账面工作表，多个工作表时必填')
    parser.add_argument('--count-sheet', help='盘点工作表，多个工作表时必填')
    parser.add_argument('--output', required=True, help='差异 CSV 输出路径')
    args = parser.parse_args()
    try:
        print(json.dumps(reconcile(args), ensure_ascii=False, indent=2))
    except (ValueError, OSError) as error:
        print(f'核对失败：{error}', file=sys.stderr)
        return 2
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
